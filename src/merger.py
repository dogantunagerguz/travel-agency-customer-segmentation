#!/usr/bin/env python3
"""
Merge reservation reports into a single flat table.

The source system exports reservations one file per month, laid out for a
person to read rather than for a machine to parse: one header row per
reservation followed by indented detail rows. This script walks a folder
(and its subfolders), turns each of those reading-oriented sheets into flat
records, and writes them out as one table.

The exports carry a travel date but no transaction date. The month and year
live in the filename instead, so the transaction date is reconstructed from
there as the first of that month. That reconstructed column is what makes
booking lead time (travel date minus transaction date) computable at all.

Usage:
    python merger.py --input <folder> --output <file.xlsx>

With no arguments it scans the current working directory.
"""

import argparse
import glob
import math
import os
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEFAULT_OUTPUT_NAME = "reservations_clean.xlsx"

# Turkish month names -> month number. Filenames come from several sources and
# the Turkish-specific characters survive inconsistently, so each month is
# mapped from its correct spelling, its ASCII fold, and the mangled form that
# appears when those characters are dropped rather than folded.
MONTHS = {
    "OCAK": 1,
    "UBAT": 2, "SUBAT": 2, "ŞUBAT": 2,
    "MART": 3,
    "NSAN": 4, "NISAN": 4, "NİSAN": 4,
    "MAYIS": 5,
    "HAZRAN": 6, "HAZIRAN": 6, "HAZİRAN": 6,
    "TEMMUZ": 7,
    "AUSTOS": 8, "AGUSTOS": 8, "AĞUSTOS": 8,
    "EYLL": 9, "EYLUL": 9, "EYLÜL": 9,
    "EKM": 10, "EKIM": 10, "EKİM": 10,
    "KASIM": 11,
    "ARALIK": 12,
}

# Output columns. Names stay in Turkish because they carry through to the
# Power BI model, which reads this file directly.
COLUMNS = [
    "No", "Rez.No", "Kabul", "Ad Soyad", "Yaş", "Otel", "Geliş", "Ayrılış",
    "Gece", "İşlem Tarihi", "Toplam", "Ödenen", "Bakiye", "Yetişkin",
    "Çocuk", "Düşünceler",
]


def transaction_date_from_filename(path):
    """Reconstruct the transaction date from the filename.

    Returns the first of the month. Word order doesn't matter: both
    "2025 OCAK" and "OCAK 2026" resolve. Returns None if either the year
    or the month is missing, which the caller reports rather than guesses.
    """
    stem = os.path.splitext(os.path.basename(path))[0].upper()
    tokens = re.split(r"[^0-9A-ZÇĞİÖŞÜ]+", stem)
    year = month = None
    for token in tokens:
        if re.fullmatch(r"20\d{2}", token):
            year = int(token)
        elif token in MONTHS:
            month = MONTHS[token]
    if year and month:
        return datetime(year, month, 1).date()
    return None


def is_number(value):
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    return isinstance(value, str) and value.strip().isdigit()


def parse_money(value):
    """Parse a currency cell into a float.

    Turkish formatting uses '.' for thousands and ',' for decimals, but the
    exports aren't consistent, so the separator role is inferred from
    position rather than assumed. Preserve minus signs and accounting-style
    parentheses so refunds and negative balances keep their sign.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return round(value, 2) if math.isfinite(value) else None

    raw = str(value).strip().replace("\u2212", "-")
    accounting_negative = raw.startswith("(") and raw.endswith(")")
    if accounting_negative:
        raw = raw[1:-1].strip()
    text = re.sub(r"[^\d.,+-]", "", raw)
    if not text:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        # A comma followed by one or two digits is a decimal; otherwise it's
        # a thousands separator.
        text = text.replace(",", ".") if re.search(r",\d{1,2}$", text) else text.replace(",", "")
    try:
        amount = float(text)
        if not math.isfinite(amount):
            return None
        return round(-abs(amount) if accounting_negative else amount, 2)
    except ValueError:
        return None


def parse_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    text = str(value).strip()
    for fmt in ("%d-%m-%Y", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def split_name_age(value):
    """Names arrive with the age appended in brackets: 'Jane Doe [34]'."""
    if value is None:
        return ("", None)
    text = str(value)
    match = re.search(r"\[(\d+)\]", text)
    name = re.sub(r"\s*\[\d+\]\s*", "", text).strip()
    return (name, int(match.group(1)) if match else None)


def parse_occupancy(value):
    """Occupancy arrives as free text, e.g. '2 Yetişkin 1 Çocuk'."""
    text = str(value or "")
    adults = re.search(r"(\d+)\s*Yeti", text)
    children = re.search(r"(\d+)\s*Çocuk", text)
    return (
        int(adults.group(1)) if adults else None,
        int(children.group(1)) if children else 0,
    )


def first_value(row, start=2):
    """Detail rows put their value in an unpredictable column."""
    for value in row[start:]:
        if value not in (None, ""):
            return value
    return None


def extract_records(worksheet):
    """Flatten one reading-oriented sheet into reservation records.

    A row starting with a number opens a new reservation. The rows beneath it
    are detail lines belonging to that reservation, identified by their label
    rather than their position.
    """
    records, current = [], None
    for raw in worksheet.iter_rows(values_only=True):
        row = list(raw) + [None] * (13 - len(raw))
        first = row[0]

        if is_number(first) and row[1] not in (None, ""):
            name, age = split_name_age(row[3])
            arrival, departure = parse_date(row[5]), parse_date(row[6])
            current = {
                "No": int(float(first)),
                "Rez.No": str(row[1]).strip(),
                "Kabul": row[2],
                "Ad Soyad": name,
                "Yaş": age,
                "Otel": row[4],
                "Geliş": arrival,
                "Ayrılış": departure,
                "Gece": (departure - arrival).days if (arrival and departure) else None,
                "İşlem Tarihi": None,  # filled in from the filename by the caller
                "Toplam": parse_money(row[7]),
                "Ödenen": parse_money(row[8]),
                "Bakiye": parse_money(row[9]),
                "Yetişkin": None,
                "Çocuk": 0,
                "Düşünceler": "",
            }
            records.append(current)

        elif current is not None and isinstance(first, str):
            if "Kişi Adedi" in first:
                current["Yetişkin"], current["Çocuk"] = parse_occupancy(first_value(row))
            elif "Düşünceler" in first:
                note = first_value(row)
                current["Düşünceler"] = str(note).strip() if note else ""

    return records


def write_output(records, path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Rezervasyonlar"
    worksheet.append(COLUMNS)
    for record in records:
        worksheet.append([record.get(column) for column in COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for column in range(1, len(COLUMNS) + 1):
        cell = worksheet.cell(row=1, column=column)
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    money_columns = {COLUMNS.index(c) + 1 for c in ("Toplam", "Ödenen", "Bakiye")}
    date_columns = {COLUMNS.index(c) + 1 for c in ("Geliş", "Ayrılış", "İşlem Tarihi")}
    text_columns = {COLUMNS.index("Rez.No") + 1}

    for row in range(2, worksheet.max_row + 1):
        for column in range(1, len(COLUMNS) + 1):
            cell = worksheet.cell(row=row, column=column)
            cell.font, cell.border = cell_font, border
            if column in money_columns:
                cell.number_format = '#,##0.00" ₺"'
            elif column in date_columns:
                cell.number_format = "DD-MM-YYYY"
            elif column in text_columns:
                # Reservation numbers are identifiers, not quantities. Excel
                # strips their leading zeros unless the cell is text.
                cell.number_format = "@"
        if row % 2 == 0:
            for column in range(1, len(COLUMNS) + 1):
                worksheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor="F2F6FA")

    widths = [5, 14, 10, 22, 6, 30, 12, 12, 6, 13, 13, 13, 11, 9, 7, 60]
    for index, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{worksheet.max_row}"
    workbook.save(path)


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input", default=os.getcwd(),
        help="Folder containing the monthly exports. Subfolders are scanned too. "
             "Defaults to the current directory.",
    )
    parser.add_argument(
        "--output", default=None,
        help=f"Output file. Defaults to {DEFAULT_OUTPUT_NAME} inside the input folder.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    input_dir = args.input
    output_path = args.output or os.path.join(input_dir, DEFAULT_OUTPUT_NAME)
    output_name = os.path.basename(output_path)

    files = [
        path for path in glob.glob(os.path.join(input_dir, "**", "*.xlsx"), recursive=True)
        if output_name not in os.path.basename(path)
        and not os.path.basename(path).startswith("~$")  # Excel lock files
    ]
    print(f"{len(files)} file(s) found\n")

    records, skipped, undated = [], 0, []
    for path in files:
        try:
            workbook = load_workbook(path)
            found = extract_records(workbook[workbook.sheetnames[0]])
            transaction_date = transaction_date_from_filename(path)
            if transaction_date is None:
                undated.append(os.path.basename(path))
            for record in found:
                record["İşlem Tarihi"] = transaction_date
            records += found
            label = transaction_date.strftime("%d.%m.%Y") if transaction_date else "NO DATE FOUND"
            print(f"  OK       {os.path.basename(path)}  (+{len(found)})  [{label}]")
        except Exception as error:
            # A single unreadable file shouldn't lose the whole run.
            skipped += 1
            print(f"  SKIPPED  {os.path.basename(path)}  ->  {error}")

    if not records:
        print(
            "\nNo records extracted. Files listed as SKIPPED could not be opened. "
            "If nothing was listed at all, check the --input path."
        )
        return

    write_output(records, output_path)
    print(f"\nDone: {len(records)} reservations, {skipped} file(s) skipped.")
    if undated:
        print(
            f"WARNING: month/year unreadable in {len(undated)} file(s), "
            f"transaction date left blank: {undated}"
        )
    print(f"Output: {output_path}")


if __name__ == "__main__":
    main()
