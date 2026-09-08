#!/usr/bin/env python3
"""Load the Power BI demo workbooks, execute SQL, and render reviewable results."""

import argparse
from datetime import date, datetime
import importlib.util
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

SQL_DIR = Path(__file__).resolve().parent
ROOT = SQL_DIR.parent
SQL_FILES = [
    "01_raw_schema.sql",
    "02_staging.sql",
    "03_marts.sql",
    "04_quality_checks.sql",
]
REPORTS = [
    ("Pipeline reconciliation", "SELECT * FROM pipeline_reconciliation"),
    ("Lifecycle distribution", """
        SELECT lifecycle_segment, COUNT(*) AS customers
        FROM customer_lifecycle_2026
        GROUP BY lifecycle_segment
        ORDER BY lifecycle_segment
    """),
    ("Top customer value records", """
        SELECT customer_name, booking_frequency, lifetime_value, value_rank
        FROM customer_value_2026
        ORDER BY value_rank, customer_name LIMIT 5
    """),
    ("2026 monthly KPIs", """
        SELECT * FROM monthly_booking_kpis
        WHERE booking_month LIKE '2026-%'
        ORDER BY booking_month
    """),
    ("Data-quality checks", """
        SELECT * FROM quality_check_results ORDER BY check_name
    """),
]


def load_setup_module():
    path = ROOT / "scripts/setup_demo.py"
    spec = importlib.util.spec_from_file_location("travel_setup_demo", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def rows_from_workbook(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers = next(values)
        return [dict(zip(headers, row)) for row in values]
    finally:
        workbook.close()


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def load_workbooks(connection, folder):
    hotel_rows = rows_from_workbook(folder / "Hotels-Mock.xlsx")
    connection.executemany("INSERT INTO raw_hotels VALUES (?, ?, ?, ?, ?, ?, ?)", [(
        row["Otel Adı"], row["Ülke"], row["Şehir"], row["İlçe"],
        row["Otel Sınıfı"], row["Enlem"], row["Boylam"]
    ) for row in hotel_rows])

    reservation_rows = rows_from_workbook(folder / "Travel-Mock.xlsx")
    connection.executemany("""
        INSERT INTO raw_reservations VALUES (
            ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
        )
    """, [(
        row["No"], row["Rez.No"], row["Kabul"], row["Ad Soyad"], row["Yaş"],
        row["Otel"], row["Geliş"], row["Ayrılış"], row["Gece"],
        iso_date(row["İşlem Tarihi"]), row["Toplam"], row["Ödenen"],
        row["Bakiye"], row["Yetişkin"], row["Çocuk"], row["Düşünceler"]
    ) for row in reservation_rows])


def build_connection():
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    connection.executescript((SQL_DIR / SQL_FILES[0]).read_text(encoding="utf-8"))
    with TemporaryDirectory() as directory:
        folder = Path(directory)
        load_setup_module().build_data(folder)
        load_workbooks(connection, folder)
    for filename in SQL_FILES[1:]:
        connection.executescript((SQL_DIR / filename).read_text(encoding="utf-8"))
    return connection


def markdown_table(rows):
    if not rows:
        return "_No rows._"
    columns = rows[0].keys()
    lines = ["| " + " | ".join(columns) + " |", "|" + "---|" * len(columns)]
    for row in rows:
        values = ["" if row[column] is None else str(row[column]) for column in columns]
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def render_results(connection):
    sections = [
        "# Executed SQL results",
        "",
        "> Generated from the same fully synthetic Excel workbooks used by the public Power BI demo. These are portfolio-demo outputs, not private or operational results.",
        "",
        "**Reference model:** fixed 2026 lifecycle windows; SQLite in-memory execution.",
    ]
    for title, query in REPORTS:
        sections.extend(["", f"## {title}", "", markdown_table(connection.execute(query).fetchall())])
    return "\n".join(sections) + "\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--write-results", action="store_true",
                        help="Refresh sql/RESULTS.md from the executed queries")
    args = parser.parse_args()
    connection = build_connection()
    try:
        output = render_results(connection)
        print(output, end="")
        failures = connection.execute(
            "SELECT check_name, issue_count FROM quality_check_results WHERE issue_count <> 0"
        ).fetchall()
        if failures:
            raise SystemExit("SQL demo failed one or more data-quality checks.")
        if args.write_results:
            (SQL_DIR / "RESULTS.md").write_text(output, encoding="utf-8")
            print("Updated sql/RESULTS.md.")
        print("PASS: all SQL data-quality checks returned zero issues.")
    finally:
        connection.close()


if __name__ == "__main__":
    main()
