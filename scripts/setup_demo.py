#!/usr/bin/env python3
"""Build a wholly synthetic, offline demo and configure its Power BI sources."""

import argparse
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
MODEL = ROOT / 'Travel-Agency-Mock.SemanticModel/definition'
PROJECT = 'Travel-Agency-Mock.pbip'
MARKER = "demo-manifest.json"


def write_book(folder, filename, sheet, columns, rows, date_columns=()):
    """Write sample rows with the exact workbook/sheet names used by the model."""
    workbook = Workbook()
    workbook.properties.creator = "Synthetic portfolio demo"
    workbook.properties.description = "Invented sample data; not business results."
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(columns)
    for row in rows:
        if len(row) != len(columns):
            raise ValueError(f"Wrong column count in {filename}")
        worksheet.append(row)
    for column in date_columns:
        index = columns.index(column) + 1
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row, index).number_format = "yyyy-mm-dd"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(folder / filename)
    workbook.close()
    return {"file": filename, "sheet": sheet, "rows": len(rows), "columns": columns}


def configure_model(model, folder):
    expressions = model / "expressions.tmdl"
    source = expressions.read_text(encoding="utf-8")
    # Power Query text literals escape quotes by doubling them.
    value = folder.resolve().as_posix().replace('"', '""')
    replacement = ('expression DemoDataFolder = "' + value
                   + '" meta [IsParameterQuery=true, Type="Text", IsParameterQueryRequired=true]')
    source, count = re.subn(r'^expression DemoDataFolder = .*$',
                            lambda match: replacement, source, flags=re.MULTILINE)
    if count != 1:
        raise ValueError("Expected one DemoDataFolder parameter in expressions.tmdl")
    expressions.write_text(source, encoding="utf-8")


def build_data(folder):
    hotel_columns = ["Otel Adı", "Ülke", "Şehir", "İlçe", "Otel Sınıfı"]
    hotels = [[f"DEMO HOTEL {i:02}", "Türkiye", city, district, "Demo Class"]
              for i, (city, district) in enumerate([
                  ("Antalya", "Manavgat"), ("Muğla", "Bodrum"), ("İzmir", "Çeşme"),
                  ("Antalya", "Alanya"), ("Aydın", "Kuşadası"), ("Muğla", "Fethiye")], 1)]
    coordinates = [[36.70 + i * .13, 27.30 + i * .31] for i in range(len(hotels))]
    files = [write_book(folder, "Hotels-Mock.xlsx", "Oteller",
                        hotel_columns + ["Enlem", "Boylam"],
                        [row + coords for row, coords in zip(hotels, coordinates)]),
             write_book(folder, "Etstur_Otel_Sinif.xlsx", "Oteller", hotel_columns, hotels)]
    price_columns = ["Oda Sınıfı", "Ort. 2 Kişilik Oda Fiyatı (₺)",
                     "Ort. 3 Kişilik Oda Fiyatı (₺)", "Ort. 4 Kişilik Oda Fiyatı (₺)",
                     "Ort. Tek Kişilik Oda Fiyatı (₺)", "Acente"]
    files.append(write_book(folder, "Didimtur_Otel_Fiyatlari.xlsx", "Oteller",
                            hotel_columns + price_columns,
                            [row + ["Demo Room", 2000, 2800, 3500, 1400, "Demo Agency"] for row in hotels]))
    columns = ["No", "Rez.No", "Kabul", "Ad Soyad", "Yaş", "Otel", "Geliş", "Ayrılış",
               "Gece", "İşlem Tarihi", "Toplam", "Ödenen", "Bakiye", "Yetişkin", "Çocuk", "Düşünceler"]
    # Six customers per lifecycle state, with repeat booking history across years.
    patterns = [(2026,), (2025, 2026), (2024,), (2023, 2024), (2023, 2026)]
    rows = []
    for customer in range(1, 31):
        for year in patterns[(customer - 1) % len(patterns)]:
            month = 1 + (customer * 3 + year) % 12
            purchased = date(year, month, 1)
            arrival = purchased + timedelta(days=20 + customer % 90)
            nights = 3 + customer % 6
            adults, children = 1 + customer % 2, customer % 3
            amount = nights * (900 + customer * 40)
            number = len(rows) + 1
            # Excel serials match the existing Int64 -> date source conversion.
            serial = lambda value: (value - date(1899, 12, 30)).days
            rows.append([number, 100000 + number, "Accepted", f"DEMO CUSTOMER {customer:03}",
                         20 + customer % 45, hotels[(customer - 1) % len(hotels)][0],
                         serial(arrival), serial(arrival + timedelta(days=nights)), nights,
                         purchased, amount, amount, 0, adults, children, "Synthetic demo"])
    files.append(write_book(folder, "Travel-Mock.xlsx", "Rezervasyonlar", columns, rows,
                            date_columns=["İşlem Tarihi"]))
    return files



def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=ROOT / "demo-data",
                        help="Output folder for generated workbooks (default: repo/demo-data)")
    args = parser.parse_args()
    folder = args.data_dir.expanduser().resolve()
    marker = folder / MARKER
    if folder.exists() and any(folder.glob("*.xlsx")):
        if not marker.exists():
            parser.error("Output contains existing workbooks; choose a new --data-dir.")
        if json.loads(marker.read_text(encoding="utf-8")).get("generator") != PROJECT:
            parser.error("Output belongs to a different demo; choose a new --data-dir.")
    folder.mkdir(parents=True, exist_ok=True)
    files = build_data(folder)
    marker.write_text(json.dumps({
        "generator": PROJECT,
        "data_kind": "fully synthetic; no private inputs or API calls",
        "reference_year": 2026,
        "files": files,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    configure_model(MODEL, folder)
    print(f"Created {len(files)} synthetic workbooks in {folder}")
    print(f"Configured DemoDataFolder. Open {ROOT / PROJECT} in Power BI Desktop and Refresh.")


if __name__ == "__main__":
    main()
