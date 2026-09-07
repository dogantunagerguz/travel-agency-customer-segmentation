"""Check a fresh, relocated public demo without private inputs or API access."""

import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


class PortableDemoTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        cls.source = Path(__file__).resolve().parents[1]
        cls.repo = Path(cls.temp.name) / "İstanbul demo with spaces"
        shutil.copytree(cls.source, cls.repo,
                        ignore=shutil.ignore_patterns(".git", "__pycache__", "demo-data"))
        cls.command = [sys.executable, str(cls.repo / "scripts/setup_demo.py")]
        cls.run_setup()
        cls.manifest = json.loads((cls.repo / "demo-data/demo-manifest.json").read_text(encoding="utf-8"))
        cls.expressions = next(cls.repo.rglob("expressions.tmdl"))
        cls.model = cls.expressions.parent

    @classmethod
    def run_setup(cls, *args, check=True):
        return subprocess.run(cls.command + list(args), cwd=cls.temp.name,
                              text=True, capture_output=True, check=check)

    def read_rows(self, filename):
        book = load_workbook(self.repo / "demo-data" / filename, data_only=True)
        try:
            rows = list(book.active.iter_rows(values_only=True))
            return [dict(zip(rows[0], row)) for row in rows[1:]]
        finally:
            book.close()

    def test_every_model_workbook_and_sheet_exists(self):
        references = []
        for path in self.model.rglob("*.tmdl"):
            text = path.read_text(encoding="utf-8")
            matches = list(re.finditer(r'File\.Contents\(DemoDataFolder & "/([^"\n]+)"\)', text))
            self.assertEqual(len(matches), text.count("File.Contents("), path.name)
            for match in matches:
                filename = match.group(1)
                tail = text[match.end():]
                sheet = re.search(r'Item="([^"]+)",Kind="Sheet"', tail).group(1)
                book = load_workbook(self.repo / "demo-data" / filename, data_only=True)
                try:
                    self.assertIn(sheet, book.sheetnames, filename)
                    self.assertGreater(book[sheet].max_row, 1, filename)
                    headers = {cell.value for cell in book[sheet][1]}
                    # Check typed source columns before any derived column is added.
                    source_steps = tail.split("Table.AddColumn", 1)[0].split("\nexpression ", 1)[0]
                    required = set(re.findall(r'\{"([^"]+)", (?:Int64\.Type|type \w+)\}', source_steps))
                    required = {name for name in required if not re.fullmatch(r"Column\d+", name)}
                    self.assertTrue(required <= headers, (filename, required - headers))
                    references.append(filename)
                finally:
                    book.close()
        self.assertEqual(set(references), {item["file"] for item in self.manifest["files"]})

    def test_parameter_points_to_relocated_demo_folder(self):
        text = self.expressions.read_text(encoding="utf-8")
        match = re.search(r'^expression DemoDataFolder = "((?:""|[^"])*)"', text, re.MULTILINE)
        self.assertIsNotNone(match)
        self.assertEqual(match.group(1).replace('""', '"'), (self.repo / "demo-data").as_posix())

    def test_foreign_keys_and_dimension_keys_match(self):
        filenames = {item["file"] for item in self.manifest["files"]}
        if "Travel-Mock.xlsx" in filenames:
            hotels = self.read_rows("Hotels-Mock.xlsx")
            keys = [row["Otel Adı"] for row in hotels]
            self.assertEqual(len(keys), len(set(keys)))
            self.assertTrue({row["Otel"] for row in self.read_rows("Travel-Mock.xlsx")} <= set(keys))
            self.assertTrue(all(row["Enlem"] is not None and row["Boylam"] is not None for row in hotels))
        elif "Yıllık Kursiyer Listesi.xlsx" in filenames:
            rows = self.read_rows("Yıllık Kursiyer Listesi.xlsx")
            keys = [row["ADAY NO"] for row in rows]
            self.assertEqual(len(keys), len(set(keys)))
            for filename in ["Gelir Listesi.xlsx", "Kursiyer Genel Sınav ve Borç Listesi.xlsx"]:
                self.assertEqual({row["ADAY NO"] for row in self.read_rows(filename)}, set(keys))
        else:
            trainees = self.read_rows("2010-2025 TARAMA.xlsx")
            results = self.read_rows("22.06.2026-son2507.xlsx")
            keys = [row["ADAY NO"] for row in results]
            self.assertEqual(len(keys), len(set(keys)))
            self.assertEqual({row["ADAY NO"] for row in trainees}, set(keys))
            self.assertTrue(all(isinstance(key, int) for key in keys))
            pools = self.read_rows("BI-Psiko 2025.xlsx") + self.read_rows("BI-Psiko 2026.xlsx")
            self.assertEqual(len(pools), len({row["Customer"] for row in pools}))

    def test_dax_field_parameters_reference_existing_columns(self):
        sources = [path.read_text(encoding="utf-8") for path in (self.model / "tables").glob("*.tmdl")]
        columns = {}
        for source in sources:
            table = source.splitlines()[0].removeprefix("table ").strip("'")
            columns[table] = set(re.findall(r"^\tcolumn (?:'([^']+)'|([^\s=]+))", source, re.MULTILINE))
            columns[table] = {first or second for first, second in columns[table]}
        for source in sources:
            for table, column in re.findall(r"NAMEOF\('([^']+)'\[([^]]+)\]\)", source):
                self.assertIn(column, columns[table], (table, column))

    def test_rerun_has_identical_sample_values(self):
        before = {item["file"]: self.read_rows(item["file"]) for item in self.manifest["files"]}
        self.run_setup()
        after = {item["file"]: self.read_rows(item["file"]) for item in self.manifest["files"]}
        self.assertEqual(before, after)

    def test_existing_non_demo_workbook_is_preserved(self):
        output = Path(self.temp.name) / "existing data"
        output.mkdir()
        sentinel = output / "keep.xlsx"
        sentinel.write_bytes(b"existing input")
        result = self.run_setup("--data-dir", str(output), check=False)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(sentinel.read_bytes(), b"existing input")


if __name__ == "__main__":
    unittest.main()
