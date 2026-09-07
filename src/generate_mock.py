#!/usr/bin/env python3
"""
Create an anonymized copy of the reservations and hotels workbooks for a public
portfolio, without breaking the model.

The two tables are joined on hotel name, and a customer's history is spread
across multiple reservation rows. So names can't be replaced row by row: the
same real hotel has to map to the same fake hotel everywhere it appears, in
both files, or the join silently drops rows and per-customer history falls
apart. This builds one mapping per distinct value and applies it consistently.

Hotel names are mapped on a NORMALIZED key (case and Turkish characters
folded), so "Güral" and "GÜRAL" collapse to the same fake hotel rather than
two. That mirrors the normalization the Power BI model already does, so the
join behaves the same on masked data as on real data.

Dates are left as-is. Lead time (the gap between transaction and travel date)
is what the report's findings depend on, and a real date next to a fake name
identifies no one. Prices are scaled by a fixed factor.

Usage:
    python generate_mock.py \
        --reservations reservations_clean.xlsx \
        --hotels Oteller.xlsx \
        --out-dir mock/
"""

import argparse
import os
from pathlib import Path
import random
import unicodedata

from openpyxl import Workbook, load_workbook

# ---- Configuration -------------------------------------------------------

PRICE_FACTOR = 1.37      # every monetary value scales by this
SEED = 20260625          # fixed so re-runs produce the same anonymization

RESERVATION_HOTEL_COL = "Otel Adı"
RESERVATION_NAME_COL = "Ad Soyad"
RESERVATION_MONEY_COLS = ["Toplam", "Ödenen", "Bakiye"]

HOTEL_NAME_COL = "Otel Adı"

# --------------------------------------------------------------------------


def normalize(value):
    """Fold case and Turkish-specific characters to a stable key.

    "Güral", "GÜRAL", and "gural" all collapse to the same key, so a hotel
    written inconsistently across the two files still maps to one fake name.
    """
    if value is None:
        return ""
    text = str(value).strip().casefold()
    # Turkish dotless-i / dotted-I don't fold the way NFKD expects; handle first
    text = text.replace("ı", "i").replace("İ".casefold(), "i")
    decomposed = unicodedata.normalize("NFKD", text)
    stripped = "".join(c for c in decomposed if not unicodedata.combining(c))
    return stripped


def scale_money(value, factor):
    if value is None or value == "":
        return value
    try:
        return round(float(value) * factor, 2)
    except (TypeError, ValueError):
        return value


def read_sheet(path):
    workbook = load_workbook(path)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    headers = list(rows[0])
    data = [list(r) for r in rows[1:]]
    return headers, data


def write_sheet(path, headers, data, sheet_title):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title
    worksheet.append(headers)
    for row in data:
        worksheet.append(row)
    workbook.save(path)


def build_hotel_map(res_headers, res_data, hotel_headers, hotel_data, rng):
    """Collect every distinct hotel across both files, keyed by normalized name,
    and assign each a stable fake name."""
    res_idx = res_headers.index(RESERVATION_HOTEL_COL)
    hotel_idx = hotel_headers.index(HOTEL_NAME_COL)

    # Map normalized key -> one representative display value, so the fake name
    # is generated once per real hotel.
    seen = {}
    for row in res_data:
        raw = row[res_idx]
        if raw not in (None, ""):
            seen.setdefault(normalize(raw), raw)
    for row in hotel_data:
        raw = row[hotel_idx]
        if raw not in (None, ""):
            seen.setdefault(normalize(raw), raw)

    keys = sorted(seen.keys())
    rng.shuffle(keys)  # so fake numbering doesn't leak alphabetical order
    return {key: f"Otel {i + 1}" for i, key in enumerate(keys)}


def build_name_map(res_headers, res_data, rng):
    """Assign each distinct customer name a stable fake name."""
    name_idx = res_headers.index(RESERVATION_NAME_COL)
    seen = {}
    for row in res_data:
        raw = row[name_idx]
        if raw not in (None, ""):
            seen.setdefault(str(raw).strip(), True)
    names = sorted(seen.keys())
    rng.shuffle(names)
    return {name: f"Müşteri {i + 1}" for i, name in enumerate(names)}


def mask_reservations(headers, data, hotel_map, name_map):
    hotel_idx = headers.index(RESERVATION_HOTEL_COL)
    name_idx = headers.index(RESERVATION_NAME_COL)
    money_idxs = [headers.index(c) for c in RESERVATION_MONEY_COLS if c in headers]

    # Dates are left untouched. Lead time (travel date minus transaction date)
    # is the report's central finding, and shifting either date risks it, while
    # a real date next to a fake name identifies no one.
    for row in data:
        raw_hotel = row[hotel_idx]
        if raw_hotel not in (None, ""):
            row[hotel_idx] = hotel_map[normalize(raw_hotel)]
        raw_name = row[name_idx]
        if raw_name not in (None, ""):
            row[name_idx] = name_map[str(raw_name).strip()]
        for i in money_idxs:
            row[i] = scale_money(row[i], PRICE_FACTOR)
    return data


def mask_hotels(headers, data, hotel_map, rng):
    hotel_idx = headers.index(HOTEL_NAME_COL)

    # Real geography would let someone reverse the renamed hotel, so it's
    # replaced rather than kept. Each fake hotel gets ONE consistent fake
    # location (Otel 1 always sits in the same city/district), drawn from real
    # Turkish place names shuffled so the mapping to the real hotel is broken.
    # The map still renders and looks plausible; none of it is traceable.
    blank_cols = ["Enlem", "Boylam", "OSM_adres",
                  "OSM_il_adaylari", "OSM_ilce_adaylari", "Konum_kaynagi"]
    blank_idxs = [headers.index(c) for c in blank_cols if c in headers]

    city_idx = headers.index("Şehir") if "Şehir" in headers else None
    district_idx = headers.index("İlçe") if "İlçe" in headers else None
    country_idx = headers.index("Ülke") if "Ülke" in headers else None

    # Real coastal-holiday provinces with a district each, shuffled and assigned
    # per fake hotel so location is plausible but decoupled from the real hotel.
    place_pool = [
        ("Antalya", "Kemer"), ("Antalya", "Manavgat"), ("Antalya", "Alanya"),
        ("Muğla", "Bodrum"), ("Muğla", "Marmaris"), ("Muğla", "Fethiye"),
        ("İzmir", "Çeşme"), ("İzmir", "Kuşadası"), ("Aydın", "Didim"),
        ("Balıkesir", "Ayvalık"), ("Antalya", "Belek"), ("Muğla", "Datça"),
    ]
    rng.shuffle(place_pool)

    fake_hotels = sorted(set(hotel_map.values()))
    place_for = {h: place_pool[i % len(place_pool)] for i, h in enumerate(fake_hotels)}

    for row in data:
        raw = row[hotel_idx]
        if raw not in (None, ""):
            fake = hotel_map[normalize(raw)]
            row[hotel_idx] = fake
            city, district = place_for[fake]
            if city_idx is not None:
                row[city_idx] = city
            if district_idx is not None:
                row[district_idx] = district
            if country_idx is not None:
                row[country_idx] = "Türkiye"
        for i in blank_idxs:
            row[i] = None
    return data


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reservations", type=Path, required=True)
    parser.add_argument("--hotels", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    RESERVATIONS_PATH = args.reservations.expanduser().resolve()
    HOTELS_PATH = args.hotels.expanduser().resolve()
    OUT_DIR = args.out_dir.expanduser().resolve()
    if any(OUT_DIR / source.name == source for source in [RESERVATIONS_PATH, HOTELS_PATH]):
        parser.error("Choose an output directory different from the source directory.")
    rng = random.Random(SEED)
    os.makedirs(OUT_DIR, exist_ok=True)

    res_headers, res_data = read_sheet(RESERVATIONS_PATH)
    hotel_headers, hotel_data = read_sheet(HOTELS_PATH)

    hotel_map = build_hotel_map(res_headers, res_data, hotel_headers, hotel_data, rng)
    name_map = build_name_map(res_headers, res_data, rng)

    res_data = mask_reservations(res_headers, res_data, hotel_map, name_map)
    hotel_data = mask_hotels(hotel_headers, hotel_data, hotel_map, rng)

    res_out = os.path.join(OUT_DIR, os.path.basename(RESERVATIONS_PATH))
    hotel_out = os.path.join(OUT_DIR, os.path.basename(HOTELS_PATH))
    write_sheet(res_out, res_headers, res_data, "Rezervasyonlar")
    write_sheet(hotel_out, hotel_headers, hotel_data, "Oteller")

    print(f"{len(hotel_map)} distinct hotels, {len(name_map)} distinct customers anonymized")
    print(f"  {res_out}")
    print(f"  {hotel_out}")
    print("\nNext: point Power BI's data source settings at these files and refresh.")


if __name__ == "__main__":
    main()
