"""Regression checks for signed amounts and project-local API configuration."""

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src import geocode, merger


class MoneyParsingTests(unittest.TestCase):
    def test_negative_currency_formats(self):
        cases = [-1250.5, -1250, "-1.250,50", "-1,250.50", "₺ -1.250,50",
                 "−1.250,50", "(1.250,50)"]
        for value in cases:
            with self.subTest(value=value):
                expected = -1250.0 if value == -1250 else -1250.5
                self.assertEqual(merger.parse_money(value), expected)

    def test_positive_and_zero_values(self):
        cases = [(1250.5, 1250.5), ("1.250,50", 1250.5),
                 ("1,250.50", 1250.5), ("+1.250,50", 1250.5),
                 ("12,50", 12.5), ("12.50", 12.5), ("1250.50", 1250.5),
                 ("0.25", 0.25), ("-12.50", -12.5), (1.25, 1.25),
                 (0, 0), ("0,00", 0)]
        for value, expected in cases:
            with self.subTest(value=value):
                self.assertEqual(merger.parse_money(value), expected)

    def test_dot_grouped_whole_currency_amounts(self):
        cases = [("1.250 ₺", 1250), ("12.500 TL", 12500),
                 ("1.250.000 ₺", 1250000), ("1.250", 1250),
                 ("+1.250", 1250), ("-1.250 ₺", -1250),
                 ("−1.250 TL", -1250), ("(1.250 ₺)", -1250),
                 ("(1.250.000 ₺)", -1250000)]
        for value, expected in cases:
            with self.subTest(value=value):
                self.assertEqual(merger.parse_money(value), expected)

    def test_missing_invalid_and_nonfinite_values(self):
        for value in (None, "", "n/a", True, "1-2", "--12", "1.25.000", "12..500",
                      float("nan"), float("inf")):
            with self.subTest(value=value):
                self.assertIsNone(merger.parse_money(value))

    def test_dot_grouped_amounts_survive_excel_output(self):
        workbook = Workbook()
        self.addCleanup(workbook.close)
        workbook.active.append([
            1, "000123", "Accepted", "Test Customer [34]", "Test Hotel",
            "01.09.2026", "03.09.2026", "1.250.000 ₺", "12.500 TL", "1.237.500 TL",
        ])
        records = merger.extract_records(workbook.active)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "whole_lira.xlsx"
            merger.write_output(records, output)
            result = load_workbook(output, data_only=True)
            try:
                values = dict(zip(merger.COLUMNS, next(result.active.iter_rows(
                    min_row=2, max_row=2, values_only=True))))
                self.assertEqual(values["Toplam"], 1250000)
                self.assertEqual(values["Ödenen"], 12500)
                self.assertEqual(values["Bakiye"], 1237500)
                self.assertEqual(values["Rez.No"], "000123")
            finally:
                result.close()

    def test_negative_amounts_survive_excel_output(self):
        workbook = Workbook()
        workbook.active.append([
            1, "000123", "Accepted", "Test Customer [34]", "Test Hotel",
            "01.09.2026", "03.09.2026", "-1.250,50", -1000, "(250,50)",
        ])
        records = merger.extract_records(workbook.active)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "refund.xlsx"
            merger.write_output(records, output)
            result = load_workbook(output, data_only=True)
            try:
                values = dict(zip(merger.COLUMNS, next(result.active.iter_rows(
                    min_row=2, max_row=2, values_only=True))))
                self.assertEqual(values["Toplam"], -1250.5)
                self.assertEqual(values["Ödenen"], -1000)
                self.assertEqual(values["Bakiye"], -250.5)
                self.assertEqual(values["Rez.No"], "000123")
            finally:
                result.close()
        workbook.close()


class ApiKeyTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.env_file = self.root / ".env"
        file_patch = patch.object(geocode, "__file__", str(self.root / "src" / "geocode.py"))
        file_patch.start()
        self.addCleanup(file_patch.stop)
        env_patch = patch.dict(os.environ, {}, clear=True)
        env_patch.start()
        self.addCleanup(env_patch.stop)

    def test_loads_project_dotenv_from_another_directory(self):
        self.env_file.write_text('LOCATIONIQ_API_KEY="test-file-key" # test only\n')
        elsewhere = self.root / "elsewhere"
        elsewhere.mkdir()
        previous = Path.cwd()
        try:
            os.chdir(elsewhere)
            self.assertEqual(geocode.get_api_key(), "test-file-key")
        finally:
            os.chdir(previous)

    def test_exported_key_takes_precedence(self):
        self.env_file.write_text("LOCATIONIQ_API_KEY=test-file-key\n")
        os.environ["LOCATIONIQ_API_KEY"] = "test-exported-key"
        self.assertEqual(geocode.get_api_key(), "test-exported-key")

    def test_exported_key_works_without_dotenv(self):
        os.environ["LOCATIONIQ_API_KEY"] = "test-exported-key"
        self.assertEqual(geocode.get_api_key(), "test-exported-key")

    def test_missing_key_reports_configuration_error(self):
        with self.assertRaisesRegex(SystemExit, "LOCATIONIQ_API_KEY is not set"):
            geocode.get_api_key()

    def test_empty_dotenv_key_reports_configuration_error(self):
        self.env_file.write_text("LOCATIONIQ_API_KEY=\n")
        with self.assertRaisesRegex(SystemExit, "LOCATIONIQ_API_KEY is not set"):
            geocode.get_api_key()


if __name__ == "__main__":
    unittest.main()
